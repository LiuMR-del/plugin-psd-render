"""Excel 解析与列识别.

读第一个 sheet, 首行为表头:
- SKU 列必须存在(大小写不敏感)
- Hash 列可选: 存在且至少一行值为 8 位字母数字 -> 启用 hash 命名模式
- 交易编号列可选: 非 hash 模式下存在则启用交易编号命名模式(空值行跳过)
- 定制信息列可选: 保存每行原始「key: value」文案, 不作为普通图层列
- GUI 里选中的「变体字段」列用于挑模版, 其余所有列 = 「图层名 -> 替换文案」

命名模式优先级: hash > 交易编号 > 普通(sku)。
"""

from __future__ import annotations

from dataclasses import dataclass, field

_ILLEGAL_FILENAME_CHARS = '\\/:*?"<>|'
CUSTOM_INFO_HEADERS = {"定制信息", "定制文案"}


@dataclass
class ExcelRow:
    row_no: int                     # Excel 行号(从 2 开始)
    sku: str
    hash: str = ""
    order_no: str = ""              # 交易编号, "" 表示无
    values: dict = field(default_factory=dict)  # 列名 -> 单元格文本(空单元格不收)
    custom_info: str = ""            # 定制信息列原文, "" 表示无


@dataclass
class ExcelData:
    headers: list                   # 全部表头
    custom_columns: list            # 除保留列/定制信息外的列(变体字段候选 + 图层替换列)
    rows: list                      # list[ExcelRow]
    hash_mode: bool = False
    order_mode: bool = False        # 交易编号命名模式(非 hash 模式且存在交易编号列)
    skipped: list = field(default_factory=list)  # [(行号, 原因)]
    custom_info_column: str = ""     # 实际识别到的定制信息列表头


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_valid_hash(value: str) -> bool:
    return len(value) == 8 and value.isalnum()


def parse_excel(path: str, name_column: str = "") -> ExcelData:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Excel 为空, 缺少表头行")
        headers = [_cell_to_str(c) for c in header_row]

        sku_idx = hash_idx = order_idx = custom_info_idx = None
        for i, h in enumerate(headers):
            if h.lower() == "sku" and sku_idx is None:
                sku_idx = i
            elif h.lower() == "hash" and hash_idx is None:
                hash_idx = i
            elif h == "交易编号" and order_idx is None:
                order_idx = i
            elif h.casefold() in {name.casefold() for name in CUSTOM_INFO_HEADERS} and custom_info_idx is None:
                custom_info_idx = i
        if sku_idx is None:
            raise ValueError("Excel 缺少 SKU 列(表头需包含「SKU」)")
        reserved = {sku_idx, hash_idx, order_idx, custom_info_idx}
        custom_columns = [h for i, h in enumerate(headers) if h and i not in reserved]

        rows: list[ExcelRow] = []
        skipped: list[tuple[int, str]] = []
        for row_no, cells in enumerate(rows_iter, start=2):
            texts = [_cell_to_str(c) for c in cells]
            # 补齐短行(read_only 模式尾部空单元格可能缺失)
            texts += [""] * (len(headers) - len(texts))
            if not any(texts):
                continue
            sku = texts[sku_idx]
            if not sku:
                skipped.append((row_no, "缺少 SKU"))
                continue
            values = {
                headers[i]: texts[i]
                for i in range(len(headers))
                if headers[i] and i not in reserved and texts[i]
            }
            rows.append(ExcelRow(
                row_no=row_no, sku=sku,
                hash=texts[hash_idx] if hash_idx is not None else "",
                order_no=texts[order_idx] if order_idx is not None else "",
                values=values,
                custom_info=texts[custom_info_idx] if custom_info_idx is not None else "",
            ))
    finally:
        wb.close()

    custom_name_mode = bool(str(name_column or "").strip())
    hash_mode = (
        not custom_name_mode
        and hash_idx is not None
        and any(_is_valid_hash(r.hash) for r in rows)
    )
    if hash_mode:
        kept: list[ExcelRow] = []
        seen: set[str] = set()
        for r in rows:
            if not _is_valid_hash(r.hash):
                skipped.append((r.row_no, f"Hash 无效(需 8 位字母数字): {r.hash or '空'}"))
            elif r.hash in seen:
                skipped.append((r.row_no, f"Hash 重复: {r.hash}"))
            else:
                seen.add(r.hash)
                kept.append(r)
        rows = kept

    # 交易编号命名模式: hash 优先, 仅非 hash 模式且存在交易编号列时启用; 空交易编号行跳过
    order_mode = order_idx is not None and not hash_mode and not custom_name_mode
    if order_mode:
        kept = []
        for r in rows:
            if not r.order_no:
                skipped.append((r.row_no, "缺少交易编号"))
            else:
                kept.append(r)
        rows = kept

    skipped.sort(key=lambda item: item[0])
    return ExcelData(
        headers=headers, custom_columns=custom_columns,
        rows=rows, hash_mode=hash_mode, order_mode=order_mode, skipped=skipped,
        custom_info_column=headers[custom_info_idx] if custom_info_idx is not None else "",
    )


def sanitize_filename(name: str) -> str:
    # 控制字符(0-31)在 Windows 文件名里非法, 写文件会报 WinError 123
    return "".join("_" if c in _ILLEGAL_FILENAME_CHARS or ord(c) < 32 else c for c in name)


def output_name_key(name: str) -> str:
    """输出文件名的重名判断键: Windows 和 macOS 默认文件系统不区分大小写。"""
    return str(name or "").casefold()


def row_column_value(row: ExcelRow, column: str) -> str:
    """Read a configured naming column, including reserved Excel columns."""
    key = str(column or "").strip()
    if not key:
        return ""
    if key.casefold() == "sku":
        return row.sku
    if key.casefold() == "hash":
        return row.hash
    if key == "交易编号":
        return row.order_no
    return str(row.values.get(key, "") or "").strip()


def output_basename(
    row: ExcelRow, hash_mode: bool, variant_field: str, seq: int,
    name_column: str = "", template_count: int = 1,
) -> str:
    """输出文件名(兼容 SCM 预处理上传的 {8位hash}-{index} 匹配规则)。

    hash 模式:      {hash}-{序号}.png
    交易编号模式:   {交易编号}.png
    普通模式:      {sku}[-字段值]-{行号}-{序号}.png

    配置 name_column 时优先使用该列；单模板直接使用列值，多模板追加序号。
    name_column 留空时完全沿用原命名模式。
    """
    name_column = str(name_column or "").strip()
    if name_column:
        raw_name = row_column_value(row, name_column)
        if not raw_name:
            raise ValueError(f"图片命名列「{name_column}」为空")
        if raw_name.lower().endswith(".png"):
            raw_name = raw_name[:-4]
        base = sanitize_filename(raw_name).strip().rstrip(".")
        if not base:
            raise ValueError(f"图片命名列「{name_column}」无法生成有效文件名")
        suffix = f"-{seq}" if int(template_count or 1) > 1 else ""
        return f"{base}{suffix}.png"
    if hash_mode:
        return f"{row.hash}-{seq}.png"
    if row.order_no:
        return sanitize_filename(row.order_no) + ".png"
    parts = [row.sku]
    if variant_field and row.values.get(variant_field):
        parts.append(row.values[variant_field])
    parts.append(str(row.row_no))
    parts.append(str(seq))
    return sanitize_filename("-".join(parts)) + ".png"
