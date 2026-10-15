"""Output folders by date/shop/category and yellow rows in the order workbook (2026-10-15)."""
import hashlib
import os
import tempfile
import unittest
import zipfile
from datetime import date

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from PIL import Image

import excel_data
import excel_marker
import output_layout


class LayoutTests(unittest.TestCase):
    def test_category_from_excel_name(self):
        for name, expected in (("10.15毛毯.xlsx", "毛毯"), ("/x/10.12铝板画.xlsx", "铝板画"),
                               ("2026-10-15 相框.xlsx", "相框"), ("10月15日-毛毯.xlsx", "毛毯"),
                               ("毛毯订单.xlsx", "毛毯订单"), ("10.15.xlsx", "未分类")):
            self.assertEqual(output_layout.category_from_excel(name), expected, name)

    def test_row_dir(self):
        row = excel_data.ExcelRow(row_no=2, sku="S", values={"店铺名": "示例店铺-us"})
        self.assertEqual(output_layout.row_dir("/out", row, "毛毯", date(2026, 10, 15)),
                         os.path.join("/out", "10.15", "示例店铺-us", "毛毯"))
        blank = excel_data.ExcelRow(row_no=3, sku="S", values={"店铺名": " "})
        self.assertEqual(output_layout.shop_of(blank), "未知店铺")
        odd = excel_data.ExcelRow(row_no=4, sku="S", values={"店铺名": "a/b:c"})
        self.assertEqual(output_layout.shop_of(odd), "a_b_c")
        self.assertEqual(output_layout.folder_name("..", "未分类"), "未分类")
        self.assertEqual(output_layout.date_folder(date(2026, 1, 5)), "1.5")


class MarkerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.tmp.name, "10.15毛毯.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["序号", "尺寸", "图片", "SKU", "店铺名"])
        for i in range(1, 5):
            ws.append([i, '50"x60"', None, f"SKU{i}", "店A"])
        ws["B3"].font = Font(bold=True)
        ws["A4"].fill = PatternFill("solid", fgColor="FFC6EFCE")
        picture = os.path.join(self.tmp.name, "p.png")
        Image.new("RGB", (8, 8), "red").save(picture)
        ws.add_image(XLImage(picture), "C2")
        wb.save(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def _entries(self):
        with zipfile.ZipFile(self.path) as archive:
            return {info.filename: hashlib.md5(archive.read(info)).hexdigest() for info in archive.infolist()}

    def test_marks_rows_keeps_pictures_and_backs_up(self):
        before = self._entries()
        result = excel_marker.mark_rows(self.path, [2, 3, 9])
        self.assertEqual(result["marked"], [2, 3])
        self.assertEqual(result["missing"], [9])
        self.assertTrue(os.path.exists(result["backup"]))
        self.assertEqual(os.path.dirname(result["backup"]), os.path.join(self.tmp.name, "备份"))
        after = self._entries()
        self.assertEqual(set(before), set(after))
        changed = {name for name in before if before[name] != after[name]}
        self.assertEqual(changed, {"xl/styles.xml", "xl/worksheets/sheet1.xml"})
        ws = openpyxl.load_workbook(self.path).active
        for col in range(1, 6):
            self.assertEqual(ws.cell(2, col).fill.fgColor.rgb, "FFFFFF00")
            self.assertEqual(ws.cell(3, col).fill.fgColor.rgb, "FFFFFF00")
            self.assertNotEqual(ws.cell(5, col).fill.fgColor.rgb, "FFFFFF00")
        self.assertTrue(ws["B3"].font.bold, "其他格式保留")
        self.assertEqual(ws["A4"].fill.fgColor.rgb, "FFC6EFCE", "未标记行的原底色不变")
        self.assertEqual(ws["D3"].value, "SKU2")
        self.assertEqual(len(excel_data.parse_excel(self.path).rows), 4)

    def test_marking_again_is_stable(self):
        excel_marker.mark_rows(self.path, [2])
        with zipfile.ZipFile(self.path) as archive:
            styles = archive.read("xl/styles.xml")
        second = excel_marker.mark_rows(self.path, [2, 4])
        with zipfile.ZipFile(self.path) as archive:
            self.assertEqual(archive.read("xl/styles.xml").count(b"FFFFFF00"), styles.count(b"FFFFFF00"))
        self.assertEqual(second["marked"], [2, 4])
        ws = openpyxl.load_workbook(self.path).active
        self.assertEqual(ws["A4"].fill.fgColor.rgb, "FFFFFF00")
        self.assertEqual(len(os.listdir(os.path.join(self.tmp.name, "备份"))), 2)

    def test_nothing_to_mark_leaves_file_alone(self):
        before = self._entries()
        self.assertEqual(excel_marker.mark_rows(self.path, [])["backup"], "")
        self.assertEqual(self._entries(), before)
        self.assertFalse(os.path.exists(os.path.join(self.tmp.name, "备份")))

    def test_lock_file(self):
        self.assertEqual(excel_marker.lock_file(self.path), "")
        lock = os.path.join(self.tmp.name, "~$10.15毛毯.xlsx")
        open(lock, "w").close()
        self.assertEqual(excel_marker.lock_file(self.path), lock)


if __name__ == "__main__":
    unittest.main()
